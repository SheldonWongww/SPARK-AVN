"""Fill NavTTA_benchmark_results.xlsx with results already established in the repo.

Data provenance (see the two source folders):
  * AVN   : avn/results/AVN_MAIN_COMPARISON.md  (all rows are provisional [P];
            only single-source / multi-source exist, no Noisy; FeedTTA/ATENA not run).
  * VLN   : vln/results/source_baselines_20260810.json (FORMAL Source for every
            benchmark) and vln/results/analysis/hparam_search/* (R2R val_seen TTA
            finalists reproduced by the user; TTA CSV carries only SR / SPL).

Only cells with an actual source are written. Everything else (Noisy, StreamVLN,
IDEA, OURS, REVERIE/R2R-CE TTA, AVN FeedTTA/ATENA, R2R val_unseen TTA) stays blank.

The layout (rows/cols) is derived from the generator's BENCHMARKS/METHODS so the
two scripts cannot drift apart.
"""
import os
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from gen_vln_benchmark_xlsx import BENCHMARKS, METHODS  # noqa: E402

XLSX = "/Users/bytedance/keyan/NaCoTTA/NavTTA/docs/literature/NavTTA_benchmark_results.xlsx"

# metric values keyed by [sheet][base_model][method][split] -> {metric: value}
# metric keys are the arrow-stripped names (SR, SPL, SNA, DTG, TL, NE, OSR, RGSPL).
DATA = {
    "AVN": {
        "SMT+AUDIO": {
            "Source": {
                "Single-source": {"SR": 54.15, "SPL": 29.42, "SNA": 42.97, "DTG": 5.00},
                "Multi-source": {"SR": 25.90, "SPL": 13.42, "SNA": 19.18, "DTG": 7.74},
            },
            "Tent": {
                "Single-source": {"SR": 55.80, "SPL": 29.82, "SNA": 43.96, "DTG": 4.82},
                "Multi-source": {"SR": 26.10, "SPL": 13.06, "SNA": 19.03, "DTG": 7.72},
            },
            "FSTTA": {
                "Single-source": {"SR": 56.55, "SPL": 30.30, "SNA": 44.89, "DTG": 4.73},
                "Multi-source": {"SR": 26.25, "SPL": 13.52, "SNA": 19.26, "DTG": 7.65},
            },
            "EAM": {
                "Single-source": {"SR": 56.20, "SPL": 30.58, "SNA": 45.21, "DTG": 4.83},
                "Multi-source": {"SR": 27.50, "SPL": 14.13, "SNA": 20.22, "DTG": 7.51},
            },
        },
        "ENMuS^3": {
            "Source": {
                "Single-source": {"SR": 66.55, "SPL": 36.05, "SNA": 51.18, "DTG": 3.21},
                "Multi-source": {"SR": 34.75, "SPL": 17.11, "SNA": 25.85, "DTG": 7.31},
            },
            "Tent": {
                "Single-source": {"SR": 67.35, "SPL": 36.69, "SNA": 51.24, "DTG": 3.26},
                "Multi-source": {"SR": 35.55, "SPL": 17.17, "SNA": 26.08, "DTG": 7.07},
            },
            "FSTTA": {
                "Single-source": {"SR": 68.55, "SPL": 37.38, "SNA": 52.35, "DTG": 3.12},
                "Multi-source": {"SR": 35.85, "SPL": 17.55, "SNA": 26.47, "DTG": 7.16},
            },
            "EAM": {
                "Single-source": {"SR": 68.15, "SPL": 36.91, "SNA": 51.84, "DTG": 3.26},
                "Multi-source": {"SR": 35.40, "SPL": 17.12, "SNA": 26.43, "DTG": 7.17},
            },
        },
    },
    "R2R": {
        "HAMT": {
            "Source": {
                "Val Seen": {"TL": 11.15, "NE": 2.51, "SR": 75.61, "SPL": 72.18},
                "Val Unseen": {"TL": 11.46, "NE": 3.62, "SR": 66.24, "SPL": 61.51},
            },
            # TL/NE backfilled from raw val_seen valid.txt logs matched on SR/SPL.
            "Tent": {"Val Seen": {"TL": 10.64, "NE": 2.45, "SR": 76.20, "SPL": 73.56}},
            "FSTTA": {"Val Seen": {"TL": 11.15, "NE": 2.45, "SR": 76.40, "SPL": 72.97}},
            "EAM": {"Val Seen": {"TL": 11.11, "NE": 2.52, "SR": 76.59, "SPL": 73.57}},
            "FeedTTA": {"Val Seen": {"TL": 11.14, "NE": 2.48, "SR": 76.00, "SPL": 72.66}},
            "ATENA": {"Val Seen": {"TL": 11.11, "NE": 2.36, "SR": 77.47, "SPL": 74.16}},
        },
        "DUET": {
            "Source": {
                "Val Seen": {"TL": 12.33, "NE": 2.28, "SR": 78.84, "SPL": 72.88},
                "Val Unseen": {"TL": 13.94, "NE": 3.31, "SR": 71.52, "SPL": 60.41},
            },
            "Tent": {"Val Seen": {"TL": 11.36, "NE": 2.38, "SR": 78.65, "SPL": 73.99}},
            "FSTTA": {"Val Seen": {"TL": 12.76, "NE": 2.27, "SR": 79.24, "SPL": 72.34}},
            "EAM": {"Val Seen": {"TL": 11.17, "NE": 2.30, "SR": 79.33, "SPL": 74.54}},
            "FeedTTA": {"Val Seen": {"TL": 11.85, "NE": 2.27, "SR": 79.33, "SPL": 73.88}},
            "ATENA": {"Val Seen": {"TL": 10.95, "NE": 2.20, "SR": 79.82, "SPL": 75.79}},
        },
        "GOAT": {
            "Source": {
                "Val Seen": {"TL": 11.69, "NE": 1.67, "SR": 84.82, "SPL": 80.05},
                "Val Unseen": {"TL": 13.30, "NE": 2.31, "SR": 78.12, "SPL": 67.58},
            },
            "Tent": {"Val Seen": {"TL": 11.31, "NE": 1.64, "SR": 84.92, "SPL": 80.32}},
            "FSTTA": {"Val Seen": {"TL": 11.55, "NE": 1.65, "SR": 84.92, "SPL": 80.21}},
            "EAM": {"Val Seen": {"TL": 11.02, "NE": 1.64, "SR": 84.82, "SPL": 80.65}},
            "FeedTTA": {"Val Seen": {"TL": 11.67, "NE": 1.66, "SR": 84.82, "SPL": 80.06}},
            "ATENA": {"Val Seen": {"TL": 11.57, "NE": 1.64, "SR": 84.92, "SPL": 80.24}},
        },
    },
    "REVERIE": {
        "HAMT": {
            "Source": {
                "Val Seen": {"OSR": 47.65, "SR": 43.29, "SPL": 40.19, "RGSPL": 25.18},
                "Val Unseen": {"OSR": 36.84, "SR": 32.95, "SPL": 30.20, "RGSPL": 17.28},
            }
        },
        "DUET": {
            "Source": {
                "Val Seen": {"OSR": 73.86, "SR": 71.75, "SPL": 63.94, "RGSPL": 51.14},
                "Val Unseen": {"OSR": 51.07, "SR": 46.98, "SPL": 33.73, "RGSPL": 23.03},
            }
        },
        "GOAT": {
            "Source": {
                "Val Seen": {"OSR": 82.36, "SR": 80.74, "SPL": 73.44, "RGSPL": 58.82},
                "Val Unseen": {"OSR": 57.97, "SR": 53.82, "SPL": 37.52, "RGSPL": 27.00},
            }
        },
    },
    "R2R-CE": {
        "ETPNav": {
            "Source": {
                "Val Seen": {"TL": 11.09, "NE": 3.58, "OSR": 74.42, "SR": 67.48, "SPL": 59.97},
                "Val Unseen": {"TL": 11.36, "NE": 4.83, "OSR": 62.59, "SR": 55.90, "SPL": 48.30},
            }
        },
        "BevBert": {
            "Source": {
                "Val Seen": {"TL": 13.00, "NE": 3.62, "OSR": 75.32, "SR": 68.38, "SPL": 59.92},
                "Val Unseen": {"TL": 12.92, "NE": 4.53, "OSR": 66.45, "SR": 58.24, "SPL": 48.26},
            }
        },
    },
}


def mkey(metric_header):
    """Strip arrow annotations so 'SR↑' / 'DTG↓' match the data keys."""
    return metric_header.replace("↑", "").replace("↓", "").strip()


def method_row_offset(method_name):
    """Row offset of a data key ('Source', 'Tent', ...) within a base-model block."""
    for i, m in enumerate(METHODS):
        if m == "Source" and method_name == "Source":
            return i
        if m.startswith("+") and m[1:].strip() == method_name:
            return i
    raise KeyError(method_name)


def main():
    wb = load_workbook(XLSX)
    bench_map = {b[0]: b for b in BENCHMARKS}

    filled = 0
    for sheet, base_data in DATA.items():
        _, base_models, splits, metrics, _ = bench_map[sheet]
        ws = wb[sheet]
        n_metrics = len(metrics)
        metric_pos = {mkey(m): j for j, m in enumerate(metrics)}

        for base, method_data in base_data.items():
            bi = base_models.index(base)
            for method, split_data in method_data.items():
                row = 3 + bi * len(METHODS) + method_row_offset(method)
                for split, metric_vals in split_data.items():
                    si = splits.index(split)
                    for mk, val in metric_vals.items():
                        col = 2 + si * n_metrics + metric_pos[mkey(mk)]
                        cell = ws.cell(row=row, column=col, value=val)
                        cell.number_format = "0.00"
                        filled += 1

    wb.save(XLSX)
    print("filled cells:", filled)


if __name__ == "__main__":
    main()

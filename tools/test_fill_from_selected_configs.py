"""Unit tests for the selected-config -> xlsx filler."""
from pathlib import Path
import os
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT / "tools"))

from gen_vln_benchmark_xlsx import BENCHMARKS, build_sheet  # noqa: E402
import fill_from_selected_configs as filler  # noqa: E402


def _build_workbook(path):
    wb = Workbook()
    wb.remove(wb.active)
    for name, bases, splits, metrics, has_time in BENCHMARKS:
        ws = wb.create_sheet(title=name)
        build_sheet(ws, bases, splits, metrics, has_time)
    wb.save(path)


class FillFromSelectedConfigsTest(unittest.TestCase):
    def _cell(self, ws, sheet, base_model, method_label, split_label, metric):
        _, base_models, splits, headers, _ = {b[0]: b for b in BENCHMARKS}[sheet]
        bi = base_models.index(base_model)
        row = 3 + bi * len(filler.METHODS) + filler._method_row_offset(method_label)
        si = splits.index(split_label)
        pos = {filler._mkey(h): j for j, h in enumerate(headers)}
        col = 2 + si * len(headers) + pos[filler._mkey(metric)]
        return ws.cell(row=row, column=col).value

    def test_writes_both_splits_for_r2r_winner(self):
        with tempfile.TemporaryDirectory() as directory:
            xlsx = os.path.join(directory, "wb.xlsx")
            _build_workbook(xlsx)
            document = {
                "schema": "navtta.vln_tta_selected_config.v1",
                "benchmark": "r2r", "setting": "duet-r2r", "method": "fstta",
                "winner": {
                    "parameters": {"lr_fast": 6e-4},
                    "metrics_seen": {"SR": 79.3, "SPL": 74.5, "NE": 2.3, "TL": 11.2},
                    "metrics_unseen": {"SR": 72.1, "SPL": 61.0, "NE": 3.2, "TL": 13.5},
                },
                "used_fallback": False,
            }
            filler.fill([document], xlsx)
            wb = load_workbook(xlsx)
            ws = wb["R2R"]
            self.assertAlmostEqual(
                self._cell(ws, "R2R", "DUET", "FSTTA", "Val Seen", "SR"), 79.3
            )
            self.assertAlmostEqual(
                self._cell(ws, "R2R", "DUET", "FSTTA", "Val Unseen", "SR"), 72.1
            )
            self.assertAlmostEqual(
                self._cell(ws, "R2R", "DUET", "FSTTA", "Val Unseen", "SPL"), 61.0
            )

    def test_reverie_metrics_map_to_rgspl_columns(self):
        with tempfile.TemporaryDirectory() as directory:
            xlsx = os.path.join(directory, "wb.xlsx")
            _build_workbook(xlsx)
            document = {
                "benchmark": "reverie", "setting": "goat-reverie", "method": "idea",
                "winner": {
                    "parameters": {"lr": 3e-3},
                    "metrics_seen": {"OSR": 83.0, "SR": 81.0, "SPL": 74.0, "RGSPL": 59.0},
                    "metrics_unseen": {"OSR": 58.5, "SR": 54.3, "SPL": 38.0, "RGSPL": 27.5},
                },
                "used_fallback": False,
            }
            filler.fill([document], xlsx)
            wb = load_workbook(xlsx)
            ws = wb["REVERIE"]
            self.assertAlmostEqual(
                self._cell(ws, "REVERIE", "GOAT", "IDEA", "Val Unseen", "RGSPL"), 27.5
            )

    def test_fallback_without_metrics_writes_nothing(self):
        with tempfile.TemporaryDirectory() as directory:
            xlsx = os.path.join(directory, "wb.xlsx")
            _build_workbook(xlsx)
            document = {
                "benchmark": "r2r", "setting": "hamt-r2r", "method": "tent",
                "winner": None, "used_fallback": True,
                "fallback_parameters": {"lr": 1e-4},
            }
            planned = filler.fill([document], xlsx, dry_run=True)
            self.assertEqual(planned, [])

    def test_osr_metric_absent_on_r2r_is_skipped(self):
        with tempfile.TemporaryDirectory() as directory:
            xlsx = os.path.join(directory, "wb.xlsx")
            _build_workbook(xlsx)
            document = {
                "benchmark": "r2r", "setting": "duet-r2r", "method": "tent",
                "winner": {
                    "parameters": {"lr": 1e-4},
                    "metrics_seen": {"SR": 79.0, "OSR": 88.0},
                    "metrics_unseen": {"SR": 72.0, "OSR": 80.0},
                },
                "used_fallback": False,
            }
            planned = filler.fill([document], xlsx, dry_run=True)
            # OSR is not an R2R column, so only the two SR cells are planned.
            keys = sorted(k for _, _, _, k, _ in planned)
            self.assertEqual(keys, ["SR", "SR"])


if __name__ == "__main__":
    unittest.main()
